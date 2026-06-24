from pathlib import Path

from docx import Document
import fitz
import xlsxwriter


BASE_DIR = Path(__file__).resolve().parents[1]
DOCX_PATH = BASE_DIR / "source_docs" / "BIM_list.docx"
PDF_PATH = BASE_DIR / "source_docs" / "site_plan.pdf"
OUT_DIR = BASE_DIR / "outputs"
OUT_XLSX = OUT_DIR / "BIM报价清单_8万以内.xlsx"
PLAN_IMAGE = OUT_DIR / "site_plan_page1.png"


def get_source_table():
    doc = Document(DOCX_PATH)
    table = doc.tables[0]
    rows = []
    for row in table.rows:
        rows.append([cell.text.strip().replace("\n", "\n") for cell in row.cells])
    return rows


def ensure_plan_image():
    OUT_DIR.mkdir(exist_ok=True)
    if PLAN_IMAGE.exists():
        return
    doc = fitz.open(PDF_PATH)
    page = doc[0]
    pix = page.get_pixmap(matrix=fitz.Matrix(3, 3), alpha=False)
    pix.save(PLAN_IMAGE)


def money_upper(amount):
    digits = "零壹贰叁肆伍陆柒捌玖"
    units = ["", "拾", "佰", "仟"]
    big_units = ["", "万", "亿"]
    integer = int(round(amount))
    if integer == 0:
        return "零元整"
    parts = []
    group_index = 0
    while integer:
        group = integer % 10000
        integer //= 10000
        if group:
            group_text = []
            zero = False
            for i in range(4):
                n = group % 10
                group //= 10
                if n == 0:
                    if group_text and not zero:
                        group_text.append("零")
                    zero = True
                else:
                    group_text.append(digits[n] + units[i])
                    zero = False
            parts.append("".join(reversed(group_text)).rstrip("零") + big_units[group_index])
        elif parts and not parts[-1].startswith("零"):
            parts.append("零")
        group_index += 1
    return "".join(reversed(parts)).strip("零") + "元整"


def main():
    OUT_DIR.mkdir(exist_ok=True)
    ensure_plan_image()
    source_rows = get_source_table()
    specs = {row[0]: row for row in source_rows[1:10]}

    quote_rows = [
        [
            1,
            "主体及附属结构三维模型（LOD400）",
            specs["1"][2],
            "项",
            1,
            14500,
            "按总平面图范围内主线、匝道、桥梁/结构及附属设施建模，满足BIM平台数据属性要求。",
        ],
        [
            2,
            "施工进度模型更新及竣工模型交付（LOD400）",
            specs["2"][2],
            "项",
            1,
            7200,
            "按阶段资料更新，交付竣工模型及数字化资料包。",
        ],
        [
            3,
            "临建布置及动态场地布置",
            specs["3"][2],
            "项",
            1,
            6000,
            "结合总平面图施工场地、通道、临水临电、安全文明及防护设施布置。",
        ],
        [
            4,
            "危大工程专项方案模拟",
            specs["4"][2],
            "项",
            1,
            4800,
            "按项目需求选做1项危大工程专项方案模拟。",
        ],
        [
            5,
            "交通导改模拟",
            specs["5"][2],
            "项",
            1,
            4500,
            "围绕主线、匝道及节点交通组织分阶段模拟。",
        ],
        [
            6,
            "工程进度模拟",
            specs["6"][2],
            "项",
            1,
            5800,
            "结合施工计划开展4D进度展示。",
        ],
        [
            7,
            "倾斜摄影模型",
            specs["7"][2],
            "项",
            1,
            5400,
            "预算控制版：按现有/单次采集数据处理，空域、外业像控及超范围航测费用另计。",
        ],
        [
            8,
            "BIM＋GIS实景融合",
            specs["8"][2],
            "项",
            1,
            6800,
            "轻量化融合场景，支持场布、交底、进度展示。",
        ],
        [
            9,
            "数字沙盘",
            specs["9"][2],
            "项",
            1,
            3200,
            "轻量展示版，包含工程概况、重点工况、视点切换和模型查看。",
        ],
        [
            10,
            "BIM工程师驻场/项目配合服务（月报制）",
            "BIM工程师现场跟踪、月度进度配合、模型问题协调、资料整理、培训支持；按36个月*人暂估计入。",
            "月",
            36,
            600,
            "预算口径为月度跟踪与关键节点配合，不等同于36个月全职驻场。",
        ],
    ]

    workbook = xlsxwriter.Workbook(OUT_XLSX)
    workbook.set_properties(
        {
            "title": "BIM应用报价清单",
            "subject": "总价控制在8万元以内",
            "author": "Codex",
            "comments": "Based on BIM清单.docx and 总平面图.pdf",
        }
    )

    fmt_title = workbook.add_format(
        {
            "bold": True,
            "font_size": 18,
            "font_name": "Microsoft YaHei",
            "align": "center",
            "valign": "vcenter",
        }
    )
    fmt_subtitle = workbook.add_format(
        {
            "font_size": 10,
            "font_name": "Microsoft YaHei",
            "align": "center",
            "valign": "vcenter",
            "font_color": "#666666",
        }
    )
    fmt_header = workbook.add_format(
        {
            "bold": True,
            "font_name": "Microsoft YaHei",
            "font_size": 10,
            "align": "center",
            "valign": "vcenter",
            "bg_color": "#D9EAF7",
            "border": 1,
        }
    )
    fmt_text = workbook.add_format(
        {
            "font_name": "Microsoft YaHei",
            "font_size": 10,
            "valign": "top",
            "text_wrap": True,
            "border": 1,
        }
    )
    fmt_center = workbook.add_format(
        {
            "font_name": "Microsoft YaHei",
            "font_size": 10,
            "align": "center",
            "valign": "vcenter",
            "border": 1,
        }
    )
    fmt_money = workbook.add_format(
        {
            "font_name": "Microsoft YaHei",
            "font_size": 10,
            "align": "right",
            "valign": "vcenter",
            "num_format": '#,##0.00',
            "border": 1,
        }
    )
    fmt_total_label = workbook.add_format(
        {
            "bold": True,
            "font_name": "Microsoft YaHei",
            "font_size": 11,
            "align": "right",
            "valign": "vcenter",
            "bg_color": "#FFF2CC",
            "border": 1,
        }
    )
    fmt_total_money = workbook.add_format(
        {
            "bold": True,
            "font_name": "Microsoft YaHei",
            "font_size": 11,
            "align": "right",
            "valign": "vcenter",
            "num_format": '#,##0.00',
            "bg_color": "#FFF2CC",
            "border": 1,
        }
    )
    fmt_note = workbook.add_format(
        {
            "font_name": "Microsoft YaHei",
            "font_size": 10,
            "valign": "top",
            "text_wrap": True,
        }
    )
    fmt_note_head = workbook.add_format(
        {
            "bold": True,
            "font_name": "Microsoft YaHei",
            "font_size": 11,
            "valign": "top",
            "bg_color": "#E2F0D9",
            "border": 1,
        }
    )

    quote_ws = workbook.add_worksheet("报价清单")
    quote_ws.hide_gridlines(2)
    quote_ws.set_landscape()
    quote_ws.set_paper(9)
    quote_ws.fit_to_pages(1, 0)
    quote_ws.set_margins(0.3, 0.3, 0.5, 0.5)
    quote_ws.merge_range("A1:H1", "BIM应用服务报价清单", fmt_title)
    quote_ws.merge_range(
        "A2:H2",
        "编制依据：BIM清单.docx、总平面图.pdf；报价控制目标：不超过人民币80,000元",
        fmt_subtitle,
    )
    quote_ws.merge_range("A3:H3", "单位：元；暂按含税综合价编制，最终以合同约定为准。", fmt_subtitle)

    headers = [
        "序号",
        "分部分项工程/服务名称",
        "技术规格、服务内容",
        "单位",
        "数量/周期",
        "综合单价",
        "合价",
        "备注",
    ]
    for col, header in enumerate(headers):
        quote_ws.write(4, col, header, fmt_header)

    widths = [7, 24, 58, 8, 10, 12, 12, 38]
    for idx, width in enumerate(widths):
        quote_ws.set_column(idx, idx, width)

    first_data_row = 5
    for idx, row in enumerate(quote_rows):
        r = first_data_row + idx
        quote_ws.write_number(r, 0, row[0], fmt_center)
        quote_ws.write(r, 1, row[1], fmt_text)
        quote_ws.write(r, 2, row[2], fmt_text)
        quote_ws.write(r, 3, row[3], fmt_center)
        quote_ws.write_number(r, 4, row[4], fmt_center)
        quote_ws.write_number(r, 5, row[5], fmt_money)
        quote_ws.write_formula(r, 6, f"=E{r+1}*F{r+1}", fmt_money, row[4] * row[5])
        quote_ws.write(r, 7, row[6], fmt_text)
        quote_ws.set_row(r, 78)

    total_row = first_data_row + len(quote_rows)
    quote_ws.merge_range(total_row, 0, total_row, 5, "报价合计", fmt_total_label)
    quote_ws.write_formula(
        total_row,
        6,
        f"=SUM(G{first_data_row+1}:G{total_row})",
        fmt_total_money,
        sum(row[4] * row[5] for row in quote_rows),
    )
    quote_ws.write(total_row, 7, "控制在80,000元以内", fmt_total_label)

    upper_row = total_row + 1
    quote_ws.merge_range(
        upper_row,
        0,
        upper_row,
        7,
        f"人民币大写：{money_upper(79800)}（小写：¥79,800.00）",
        fmt_total_label,
    )

    quote_ws.merge_range(
        upper_row + 2,
        0,
        upper_row + 2,
        7,
        "说明：本报价为预算控制版。若需完整UE数字孪生平台、多期航飞、超出总平面图范围的建模、36个月全职驻场或合同外专项模拟，费用另行协商。",
        fmt_note,
    )
    quote_ws.freeze_panes(5, 0)
    quote_ws.autofilter(4, 0, total_row - 1, 7)

    note_ws = workbook.add_worksheet("编制说明")
    note_ws.hide_gridlines(2)
    note_ws.set_column("A:A", 18)
    note_ws.set_column("B:B", 96)
    note_ws.merge_range("A1:B1", "编制说明与边界条件", fmt_title)
    notes = [
        ("项目判断", "总平面图显示为道路立交/匝道节点工程，报价范围按图示主线、匝道、桥梁/结构及附属设施相关BIM应用服务暂估。"),
        ("编制依据", "1. BIM清单.docx 中的BIM应用工程量清单；\n2. 总平面图.pdf 中的道路、匝道、节点及施工范围信息；\n3. 总价不超过80,000元的控制要求。"),
        ("报价口径", "本表为预算控制版综合报价，暂按含税综合价编制。数量为暂估，后续可按最终合同范围、模型深度、交付频次和现场服务强度调整。"),
        ("驻场口径", "驻场/项目配合服务按36个月月报制暂估，单价600元/月，主要覆盖月度跟踪、问题协调、关键节点线上/现场配合及月报资料整理；不按36个月全职驻场人员成本计价。"),
        ("成果交付", "LOD400施工深化模型、阶段进度/竣工模型、场布模型、危大方案模拟成果、交通导改模拟、4D进度模拟、倾斜摄影模型处理成果、BIM+GIS轻量融合场景、数字沙盘轻量展示文件及月度服务记录。"),
        ("不含内容", "空域审批及第三方测绘费用、多期无人机航飞、完整UE定制孪生平台开发、硬件设备、常驻全职人员工资、重大设计变更导致的重复建模、合同范围外专项方案模拟。"),
        ("付款建议", "建议按合同签订30%、模型初版提交30%、主要应用成果提交30%、竣工/资料移交10%的节点执行，可根据甲方流程调整。"),
    ]
    row = 2
    for key, value in notes:
        note_ws.write(row, 0, key, fmt_note_head)
        note_ws.write(row, 1, value, fmt_text)
        note_ws.set_row(row, 64)
        row += 1

    source_ws = workbook.add_worksheet("BIM清单摘录")
    source_ws.hide_gridlines(2)
    source_ws.set_column("A:A", 8)
    source_ws.set_column("B:B", 34)
    source_ws.set_column("C:C", 86)
    source_ws.set_column("D:D", 13)
    source_ws.set_column("E:E", 28)
    source_ws.merge_range("A1:E1", "BIM清单.docx 摘录", fmt_title)
    for r, row_values in enumerate(source_rows, start=2):
        for c, value in enumerate(row_values):
            source_ws.write(r - 1, c, value, fmt_header if r == 2 else fmt_text)
        source_ws.set_row(r - 1, 58 if r > 2 else 28)
    source_ws.freeze_panes(2, 0)

    image_ws = workbook.add_worksheet("总平面图参考")
    image_ws.hide_gridlines(2)
    image_ws.set_column("A:A", 18)
    image_ws.set_column("B:K", 16)
    image_ws.merge_range("A1:K1", "总平面图.pdf 参考图", fmt_title)
    image_ws.merge_range(
        "A2:K2",
        "PDF为矢量图纸，文字不可直接提取；下图为第一页渲染图，仅用于报价范围判断。",
        fmt_subtitle,
    )
    image_ws.insert_image("A4", str(PLAN_IMAGE), {"x_scale": 0.58, "y_scale": 0.58})

    workbook.close()
    print(OUT_XLSX.resolve())


if __name__ == "__main__":
    main()
